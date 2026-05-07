"""
convert_iiit_2025.py
Converts 2025_IIIT.xlsx (JoSAA 2025 IIIT cutoffs, 6 sheets) into
6 JSON files matching the schema used by the JEE Toolkit app:
  public/2025_j_IIIT_r1.json  ... public/2025_j_IIIT_r6.json

Run from the project root:
  python excel/josaa/convert_iiit_2025.py
"""

import json
import os
import openpyxl

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "2025_IIIT.xlsx")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "public")

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]

    # Determine round number from sheet name (e.g. 'j_IIIT_r3' -> 3)
    round_num = int(sheet_name.split("_r")[-1])

    colleges_map: dict[str, dict] = {}  # name -> {id, name, entries[]}
    college_id_counter = 1

    for row in ws.iter_rows(min_row=2, values_only=True):
        institute, program, quota, seat_type, gender, opening_rank, closing_rank = row

        # Skip rows with missing data
        if not institute or not program or opening_rank is None or closing_rank is None:
            continue

        institute = str(institute).strip()

        if institute not in colleges_map:
            colleges_map[institute] = {
                "id": college_id_counter,
                "name": institute,
                "entries": []
            }
            college_id_counter += 1

        colleges_map[institute]["entries"].append({
            "program": str(program).strip(),
            "quota": str(quota).strip() if quota else "",
            "seatType": str(seat_type).strip() if seat_type else "",
            "gender": str(gender).strip() if gender else "",
            "openingRank": int(opening_rank),
            "closingRank": int(closing_rank)
        })

    colleges_list = list(colleges_map.values())
    total_entries = sum(len(c["entries"]) for c in colleges_list)

    output = {
        "meta": {
            "source": f"2025_IIIT_r{round_num}",
            "counselling": "JoSAA",
            "year": 2025,
            "round": round_num,
            "type": "IIIT",
            "totalEntries": total_entries,
            "totalColleges": len(colleges_list)
        },
        "colleges": colleges_list
    }

    out_file = os.path.join(OUTPUT_DIR, f"2025_j_IIIT_r{round_num}.json")
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)

    print(f"[OK] {sheet_name}: {len(colleges_list)} colleges, {total_entries} entries -> {os.path.basename(out_file)}")

print("\nDone! All 6 IIIT JSON files written to public/")
